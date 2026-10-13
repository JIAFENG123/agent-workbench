from __future__ import annotations

import base64
import io

import pytest

from app.llm.client import LLMError
from app.rag.loading import ensure_allowed, extract_text, settings

# A one-page PDF holding the single text run "Hello PDF", generated once with pypdf and
# inlined so the tests exercise only the read path that production actually uses.
PDF_WITH_TEXT = base64.b64decode(
    "JVBERi0xLjMKJeLjz9MKMSAwIG9iago8PAovUHJvZHVjZXIgKHB5cGRmKQo+PgplbmRvYmoKMiAwIG9iago8PAov"
    "VHlwZSAvUGFnZXMKL0NvdW50IDEKL0tpZHMgWyA0IDAgUiBdCj4+CmVuZG9iagozIDAgb2JqCjw8Ci9UeXBlIC9D"
    "YXRhbG9nCi9QYWdlcyAyIDAgUgo+PgplbmRvYmoKNCAwIG9iago8PAovVHlwZSAvUGFnZQovUmVzb3VyY2VzIDw8"
    "Ci9Gb250IDw8Ci9GMSA2IDAgUgo+Pgo+PgovTWVkaWFCb3ggWyAwLjAgMC4wIDYxMiA3OTIgXQovUGFyZW50IDIg"
    "MCBSCi9Db250ZW50cyA1IDAgUgo+PgplbmRvYmoKNSAwIG9iago8PAovTGVuZ3RoIDQwCj4+CnN0cmVhbQpCVCAv"
    "RjEgMjQgVGYgNzIgNzAwIFRkIChIZWxsbyBQREYpIFRqIEVUCmVuZHN0cmVhbQplbmRvYmoKNiAwIG9iago8PAov"
    "VHlwZSAvRm9udAovU3VidHlwZSAvVHlwZTEKL0Jhc2VGb250IC9IZWx2ZXRpY2EKPj4KZW5kb2JqCnhyZWYKMCA3"
    "CjAwMDAwMDAwMDAgNjU1MzUgZiAKMDAwMDAwMDAxNSAwMDAwMCBuIAowMDAwMDAwMDU0IDAwMDAwIG4gCjAwMDAw"
    "MDAxMTMgMDAwMDAgbiAKMDAwMDAwMDE2MiAwMDAwMCBuIAowMDAwMDAwMjk0IDAwMDAwIG4gCjAwMDAwMDAzODQg"
    "MDAwMDAgbiAKdHJhaWxlcgo8PAovU2l6ZSA3Ci9Sb290IDMgMCBSCi9JbmZvIDEgMCBSCj4+CnN0YXJ0eHJlZgo0"
    "NTQKJSVFT0YK"
)


def _blank_pdf() -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _docx(*, with_text: bool = True) -> bytes:
    from docx import Document

    document = Document()
    if with_text:
        document.add_heading("Amazon 玩具规则", level=1)
        document.add_paragraph("标题里不能出现 best seller。")
        document.add_paragraph("")
        document.add_paragraph("第二段内容。")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "字段"
        table.rows[0].cells[1].text = "取值"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize(
    "filename",
    [
        "rules.txt",
        "rules.md",
        "rules.markdown",
        "rules.csv",
        "rules.json",
        "rules.pdf",
        "rules.docx",
    ],
)
def test_supported_extensions_pass(filename: str) -> None:
    ensure_allowed(filename, 1024)


def test_extension_matching_ignores_case() -> None:
    ensure_allowed("RULES.PDF", 1024)


@pytest.mark.parametrize("filename", ["rules.xlsx", "rules.ppt", "rules.zip", "screenshot.png"])
def test_formats_without_an_extractor_are_rejected(filename: str) -> None:
    with pytest.raises(LLMError, match="not supported yet"):
        ensure_allowed(filename, 1024)


def test_unknown_extension_is_rejected() -> None:
    with pytest.raises(LLMError, match="Unsupported file type"):
        ensure_allowed("rules.exe", 1024)


def test_files_over_the_size_limit_are_rejected() -> None:
    with pytest.raises(LLMError, match="larger than"):
        ensure_allowed("rules.txt", settings.max_upload_bytes + 1)


def test_gb18030_text_is_decoded() -> None:
    raw = "商品标题不能用 best seller。".encode("gb18030")
    assert extract_text("rules.txt", raw) == "商品标题不能用 best seller。"


def test_docx_paragraphs_tables_and_blank_lines_are_kept() -> None:
    text = extract_text("rules.docx", _docx())
    assert "标题里不能出现 best seller。" in text
    assert "第二段内容。" in text
    assert "字段\t取值" in text
    # The blank paragraph survives, which is what keeps paragraph-aware chunking working.
    assert "\n\n" in text


def test_pdf_text_is_extracted() -> None:
    assert "Hello PDF" in extract_text("rules.pdf", PDF_WITH_TEXT)


def test_corrupt_pdf_is_rejected() -> None:
    with pytest.raises(LLMError):
        extract_text("rules.pdf", b"%PDF-1.4 this is not a real pdf body")


def test_pdf_without_a_text_layer_is_rejected() -> None:
    with pytest.raises(LLMError, match="no text layer"):
        extract_text("scan.pdf", _blank_pdf())


def test_corrupt_docx_is_rejected() -> None:
    with pytest.raises(LLMError, match="Could not open"):
        extract_text("rules.docx", b"this is not a zip archive")


def test_docx_without_any_text_is_rejected() -> None:
    with pytest.raises(LLMError, match="No text found"):
        extract_text("empty.docx", _docx(with_text=False))
